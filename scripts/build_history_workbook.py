#!/usr/bin/env python3
"""Build docs/history/function-updates.xlsx — every function/feature update of the agent
as a filterable spreadsheet (plus CSV copy). Re-run: python3 scripts/build_history_workbook.py
Requires: pip install openpyxl
"""
from __future__ import annotations

import csv
import subprocess
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "docs" / "history"

# (date, version line, phase, category, type, feature/function, description, module/file, command, PR)
# type: Feature | Enhancement | Fix | Integration | Command | Infra | Test | Docs | Security | Rollback
U = [
    # ---------------- 21 Aug ----------------
    ("2026-08-21", "AI OS v0.2", "0 Foundation", "Core State", "Feature", "Unified StateStore", "Single writer for all mutable data: atomic writes, version conflict rejection, rotating backups (last 5), audit log data/audit.jsonl. Sheet demoted to import/export.", "engine/store.py", "", ""),
    ("2026-08-21", "AI OS v0.2", "0 Foundation", "Governance", "Feature", "Action queue + approval gate (C2)", "Follow-up drafts queued as PENDING_APPROVAL; approval needs SHA-256 hash of draft text; 48h expiry; idempotent.", "engine/approve.py, engine/render_approvals.py", "approve A-001 --hash", ""),
    ("2026-08-21", "AI OS v0.2", "0 Foundation", "Chief of Staff", "Feature", "Waiting-for rot detection", "Everything awaited derived into waiting_for; >14 days escalates to a decision issue.", "engine/chief_of_staff.py", "", ""),
    ("2026-08-21", "AI OS v0.2", "0 Foundation", "Chief of Staff", "Feature", "Daily brief + weekly review + pattern detection", "Brief engine, weekly review, pattern detector, daily inbox HTML.", "engine/chief_of_staff.py, tools/daily-inbox.html", "", ""),
    ("2026-08-21", "AI OS v0.3", "0 Foundation", "Automation", "Feature", "Manager loop (two cycles)", "Fast cycle every 15 min (waiting-for sweep) + full cycle 06:00 Riyadh (brief + dashboard), catch-up, write-on-change.", "engine/manager.py", "manager.py --loop", ""),
    ("2026-08-21", "AI OS v0.3", "0 Foundation", "Core State", "Enhancement", "waiting_for schema v2", "WAITING→OVERDUE state machine; one idempotent follow-up per overdue item with draft.", "engine/manager.py", "", ""),
    ("2026-08-21", "AI OS v0.3", "0 Foundation", "Chief of Staff", "Feature", "decision_requests", "Formal decision requests (options + deadline) for stalled projects; resolution logged with 30-day review.", "engine/manager.py", "", ""),
    ("2026-08-21", "AI OS v0.3", "0 Foundation", "Chief of Staff", "Feature", "Change detection section", "'What changed since last brief' (new / status changed / closed).", "engine/chief_of_staff.py", "", ""),
    ("2026-08-21", "AI OS v0.3.1", "0 Foundation", "Clinical / Voice", "Feature", "Inbound voice-call processor", "Transcript → caller/intent classification → summary → StateStore (contacts, callbacks, leads) + follow-up draft; never sends.", "engine/voice_call.py", "voice_call.py demo A..E | ingest", ""),
    ("2026-08-21", "AI OS v0.3.1", "0 Foundation", "Security", "Security", "Voice-call hard guards", "No diagnosis from calls; caller speech untrusted (injection rejected + logged); data minimisation; idempotency by call ID.", "engine/voice_call.py", "", ""),
    ("2026-08-21", "AI OS v0.4", "0 Foundation", "Learning", "Feature", "Adaptive teaching engine", "Plans → graded chunks with adaptive rulings (≥85/70–84/50–69/<50) → mastery map → spaced repetition 1/3/7/14/30.", "engine/learning_engine.py", "", ""),
    ("2026-08-21", "AI OS v0.4", "0 Foundation", "Learning", "Feature", "Review lifecycle", "SCHEDULED→DUE→PRESENTED→ANSWERED→SCORED→COMPLETED; brief shows max 2 reviews/day.", "engine/learning_engine.py", "", ""),
    ("2026-08-21", "AI OS v0.4.1", "0 Foundation", "Learning", "Enhancement", "ILPC methodology (Bob Pike)", "EAT / CPR / 90-20-8 rules; outline generator; first material LP-001 Lean Six Sigma.", "prompts/personal-training.md, materials/", "learning_engine.py outline LP-001", ""),
    ("2026-08-21", "AI OS v0.4.1", "0 Foundation", "Infra", "Infra", "Sheet template + bootstrap + smoke test + CI", "make_template, bootstrap_demo.sh, smoke_test.sh, CI workflow in docs/ci-workflow.yml.", "engine/make_template.py, scripts/", "", ""),
    ("2026-08-21", "AI OS v0.4.1", "0 Foundation", "Integration", "Integration", "ChatGPT & VS Code bridge", "Context export for chat + editor tasks + usage guide.", "engine/export_for_chat.py, .vscode/", "", ""),
    ("2026-08-21", "AI OS v0.4.1", "0 Foundation", "Automation", "Feature", "24/7 autostart", "Windows/macOS/Linux installers + daily heartbeat.", "autostart/", "", ""),
    ("2026-08-21", "AI OS v0.4.1", "0 Foundation", "Integration", "Integration", "Google Drive bridge (sheet import)", "Live sheet import, decision-speed plan, fix for public-request invalidation.", "engine/import_drive.py", "", ""),
    ("2026-08-21", "AI OS v0.4.1", "0 Foundation", "Knowledge", "Feature", "Sources & scientific learning bridge", "--sources mode; knowledge_sources section; import of 24 sources + 28 weak-point protocols; LP-002/LP-003.", "engine/import_drive.py", "--sources", ""),
    # ---------------- 22 Aug ----------------
    ("2026-08-22", "AI OS v0.4.x", "1 CoS Core", "Life OS", "Feature", "14 life doors + door of the day", "Life-domain model and daily focus door in the brief; real income baseline.", "engine/chief_of_staff.py", "/door", ""),
    ("2026-08-22", "AI OS v0.4.x", "1 CoS Core", "Knowledge", "Feature", "Knowledge asset registry", "Unified tracking of books/docs/tabs/files/contracts; filter UI; auto-sync in full cycle.", "engine/asset_registry.py", "", ""),
    ("2026-08-22", "AI OS v0.4.x", "1 CoS Core", "Prompts", "Feature", "4 new playbooks", "Travel, negotiation, writing, family prompt packs.", "prompts/travel.md, negotiation.md, writer.md, family.md", "", ""),
    ("2026-08-22", "AI OS v0.4.x", "1 CoS Core", "Docs", "Docs", "Agent skills inventory + routing table", "11 skills + 9 engines; later audited 19/19.", "docs/agent-skills.md, evaluation/skills-audit-2026-08-22.md", "", ""),
    ("2026-08-22", "AI OS v0.4.x", "1 CoS Core", "Telegram", "Feature", "Full Telegram integration (v0)", "Button approvals (C2), decisions, brief, capture; social triage; full-text search.", "engine/telegram_bot.py, engine/social_triage.py, engine/search.py", "/approve /decisions /brief", ""),
    ("2026-08-22", "AI OS v0.4.x", "1 CoS Core", "Chief of Staff", "Fix", "Complaint pattern fix", "Pattern for 'bad/complain' + cleanup of misclassified draft.", "engine/chief_of_staff.py", "", ""),
    ("2026-08-22", "AI OS v0.4.x", "1 CoS Core", "Telegram", "Fix", "getUpdates timeout", "Timeout longer than long-poll — removes timeout noise.", "engine/telegram_bot.py", "", ""),
    ("2026-08-22", "AI OS v0.4.x", "1 CoS Core", "Life OS", "Feature", "Skills wave 2: OKR, Health Guardian, finance E-S-B-I", "OKR engine, energy log, real finance, green Friday brief, voice capture from Telegram.", "engine/okr.py, engine/energy_log.py", "/okr, energy", ""),
    ("2026-08-22", "AI OS v0.5", "1 CoS Core", "Core", "Feature", "Orchestrator, permissions, memory, RAG", "Core orchestration with provenance, permission model, memory, retrieval.", "engine/orchestrator.py, permissions.py, memory.py, rag.py", "", "#1"),
    ("2026-08-22", "AI OS v0.5", "1 CoS Core", "Core", "Feature", "Observability, telemetry, evaluation, control center", "Telemetry & observability, evaluation harness, mobile control center, v05 cycle.", "engine/observability.py, telemetry.py, evaluate.py, control_center.py", "", "#1"),
    ("2026-08-22", "AI OS v0.5", "1 CoS Core", "Clinical / Rehab", "Feature", "RCJY rehabilitation leadership training", "Work context, service packages, CoS scenarios; morning routing prioritises hospital leadership; PHI never in memory/GitHub.", "knowledge/rcjy-*.md, training/", "", "#2"),
    ("2026-08-22", "AI OS v0.6", "1 CoS Core", "Self-improvement", "Feature", "Self-improving learning loop", "Reflection engine, behaviour model, self-review, decision quality scoring.", "engine/reflection_engine.py, behavior_model.py, self_review.py, decision_quality.py", "", "#3"),
    ("2026-08-22", "AI OS v0.7", "1 CoS Core", "Trust", "Feature", "Trust, change intelligence, reliability", "Change intelligence, trust dashboard, connector health, backup verification.", "engine/change_intelligence.py, trust_dashboard.py, connector_health.py, backup_verify.py", "", "#4"),
    ("2026-08-22", "AI OS v0.8", "1 CoS Core", "Integration", "Integration", "Live connectors + unified source sync", "Gmail, Calendar, Drive, GitHub, Telegram live connectors; unified inbox.", "engine/live_sync.py, unified_inbox.py, connectors/google_workspace.py, github_live.py, telegram_live.py", "", "#5"),
    ("2026-08-22", "AI OS v0.9", "1 CoS Core", "Knowledge", "Feature", "Master professional profile + Drive governance", "Profile YAML, Drive knowledge map, source policy, source governance engine.", "knowledge/master-professional-profile.yaml, engine/source_governance.py", "", "#6"),
    ("2026-08-22", "Telegram v1", "1 CoS Core", "Telegram", "Command", "Secure polling bot commands", "/start /help /profile /sources /selftest; owner lock via TELEGRAM_ALLOWED_CHAT_ID; security tests.", "connectors/telegram_bot.py", "/start /help /profile /sources /selftest", ""),
    ("2026-08-22", "Telegram v1", "1 CoS Core", "Infra", "Infra", "Railway deployment", "requirements.txt, Dockerfile, Bedrock dependency, Railway volume.", "Dockerfile, requirements.txt", "", ""),
    ("2026-08-22", "Telegram v1", "1 CoS Core", "AI Models", "Integration", "Claude on AWS Bedrock free-text answers", "Free text answered by Claude; conversations persisted to Sheets.", "connectors/telegram_bot.py", "/bedrock_test", ""),
    ("2026-08-22", "Telegram v1", "1 CoS Core", "Memory", "Feature", "Bounded memory + knowledge retrieval runtime", "Agent runtime with memory and retrieval; Railway volume conversation state.", "engine/agent_runtime.py", "/context_test", ""),
    ("2026-08-22", "Telegram v1", "1 CoS Core", "Voice", "Integration", "AWS Transcribe voice pipeline", "Private transcription with cleanup; Telegram voice notes transcribed.", "connectors/aws_transcribe.py", "", ""),
    ("2026-08-22", "Telegram v1", "1 CoS Core", "Google Sheets", "Integration", "Apps Script secure webhook", "Sheets logging/read-search/approved updates without JSON key; Sheets intelligence connector.", "connectors/google_sheets_webhook.gs, sheet_intelligence.py", "", ""),
    ("2026-08-22", "Telegram v1", "1 CoS Core", "AI Models", "Enhancement", "Profile grounding", "AI responses always grounded in verified professional profile; executive tabs prioritised.", "evaluation/profile-grounding-cases.json", "", ""),
    ("2026-08-22", "Telegram v1", "1 CoS Core", "Telegram", "Fix", "Repair Sheets commands + syntax validation", "Fixed Telegram Sheets commands, validated Python syntax.", "connectors/telegram_bot.py", "", ""),
    # ---------------- 24-25 Aug ----------------
    ("2026-08-24", "Rehab v1.3", "2 Clinical", "Clinical / Rehab", "Feature", "Executive brief discovery", "/brief reads bounded Sheets snapshot, diffs vs last, detects changes/blockers/decisions; updates Executive_Brief tab.", "connectors/brief_discovery.py", "/brief", "#8"),
    ("2026-08-24", "Rehab v1.3", "2 Clinical", "Clinical / Rehab", "Integration", "Rehab supervisor Google Form", "createRehabSupervisorForm + testRehabIntegration.", "connectors/rehab_supervisor_form.gs", "", "#8"),
    ("2026-08-24", "Phase 1.5", "2 Clinical", "Clinical / Rehab", "Feature", "Safe pre-visit intelligence", "Case-code only; safety screen first; URGENT/PRIORITY/ROUTINE review labels; irritability; gaps; hypotheses for clinician review only.", "engine/previsit_intelligence.py, prompts/previsit-intelligence.md, connectors/previsit_patient_form.gs", "", "#9"),
    ("2026-08-24", "Phase 1.5", "2 Clinical", "Clinical / Rehab", "Feature", "Approved pre-visit link workflow", "Pre-visit links issued only through approval.", "engine/previsit_intelligence.py", "", "#9"),
    ("2026-08-24", "Rehab v1.3", "2 Clinical", "Telegram", "Fix", "Brief availability + Telegram formatting", "Brief kept available before gateway upgrade; readable formatting.", "connectors/brief_discovery.py", "", ""),
    ("2026-08-24", "P0", "2 Clinical", "Telegram", "Security", "Switch to webhook + harden Sheets/privacy", "Polling → webhook; privacy hardening.", "connectors/telegram_webhook.py", "", "#10"),
    ("2026-08-24", "P0", "2 Clinical", "Telegram", "Fix", "Duplicated /brief + Sheets timeout", "Fixed duplicate brief and timeout path.", "connectors/telegram_webhook.py", "", "#11"),
    ("2026-08-25", "Memory", "2 Clinical", "Memory", "Fix", "Historical memory retrieval + save receipts", "Verified save receipts.", "engine/context_service.py", "/memory_status", ""),
    ("2026-08-25", "Memory", "2 Clinical", "Memory", "Feature", "Provenance-first semantic context retrieval", "Unified multi-source retrieval; durable provenance-aware memory; Arabic prefixes; contract concept expansion.", "engine/context_service.py", "/memory /confirm_memory /update_memory", ""),
    ("2026-08-25", "Calendar", "2 Clinical", "Google Calendar", "Feature", "Confirmed Calendar actions + reminder scheduler", "Create/confirm events, Telegram reminders, Arabic dual-hour phrasing, service-account access, consent docs.", "connectors/calendar_actions.py, calendar_intent.py", "", ""),
    # ---------------- 27-28 Aug ----------------
    ("2026-08-27", "Prod v0.5a", "3 Multi-model", "AI Models", "Integration", "OpenRouter primary + Bedrock fallback", "Model gateway with permanent Bedrock fallback; model-router CI.", "connectors/model_gateway.py", "", "#14"),
    ("2026-08-28", "Prod v0.6", "3 Multi-model", "AI Models", "Feature", "Multi-agent task delegation", "Delegate agents + council; routing and privacy guard tests.", "connectors/task_delegation.py", "/delegate /agents /council", "#15"),
    ("2026-08-28", "Prod v0.7", "3 Multi-model", "Google", "Integration", "Read-only Google knowledge gateway", "Knowledge read commands + diagnostics; allowlist & URL parsing tests.", "(rolled back)", "/google_access", "#16"),
    ("2026-08-28", "Prod v0.7.1", "3 Multi-model", "Google", "Fix", "Robust service-account parsing", "Hotfix.", "(rolled back)", "", "#17"),
    ("2026-08-28", "Prod v0.8", "3 Multi-model", "Google", "Integration", "Apps Script knowledge gateway (Sheets + Drive)", "Preferred in production; diagnostics in /google_access.", "(rolled back)", "/google_access", "#18-#20"),
    ("2026-08-28", "Prod v0.8.1", "3 Multi-model", "Google", "Rollback", "Rollback Google Knowledge", "Restored stable v0.6 workflow (8 rollback commits).", "connectors/", "", "#21"),
    ("2026-08-28", "Mission v0.7", "3 Multi-model", "AI Models", "Feature", "AI Team Mission Orchestrator", "Shared missions for Claude, GPT, Gemini; specialists via direct APIs; Gemini fallback hotfix; provider diagnostics.", "connectors/team_orchestrator.py, direct_specialists.py, provider_diagnostics.py", "/mission", "#22-#26"),
    ("2026-08-28", "Mission v0.8", "3 Multi-model", "AI Models", "Enhancement", "Sequential Gemini→GPT handoff", "Ordered specialist handoff.", "connectors/team_orchestrator.py", "", "#27"),
    ("2026-08-28", "Mission v0.9", "3 Multi-model", "AI Models", "Feature", "Bedrock-first lean token budgets", "Lightweight Bedrock team adapter; token-budgeted missions; de-identified guardrail; research capsule format.", "connectors/bedrock_team.py, lean_missions.py", "", "#28"),
    ("2026-08-28", "Mission v0.9.1", "3 Multi-model", "AI Models", "Enhancement", "Bedrock diagnostics + tiny probe", "Probe command for Bedrock team models.", "connectors/bedrock_team.py", "", "#29"),
    ("2026-08-28", "Mission v0.9.2", "3 Multi-model", "AI Models", "Enhancement", "Nova Micro default lean model", "Cheaper default for lean missions.", "connectors/lean_missions.py", "", "#30"),
    ("2026-08-28", "Prod v0.9.3", "3 Multi-model", "Context", "Feature", "On-demand operational context capsules", "Ops context + source status (v0.9.3a).", "connectors/ops_context.py", "", "#31-#32"),
    ("2026-08-28", "Prod v0.9.4", "3 Multi-model", "Google", "Enhancement", "Hardened Google context transport", "Mounted credential files; runtime failure reasons (v0.9.4a).", "connectors/google_credentials.py", "", "#34-#36"),
    ("2026-08-28", "Prod v0.9.5", "3 Multi-model", "Google Calendar", "Fix", "Calendar reminders in webhook mode", "Reminder worker kept alive under webhook.", "connectors/telegram_webhook_runtime.py", "", "#37"),
    ("2026-08-28", "Prod v0.9.6", "3 Multi-model", "Google Calendar", "Feature", "NL Calendar requests → guarded actions", "Natural-language routing + mobile-friendly confirmation.", "connectors/mobile_calendar_confirm.py", "", "#38-#39"),
    # ---------------- 29-31 Aug ----------------
    ("2026-08-29", "Portfolio v1", "4 Safety & Truth", "Portfolio", "Feature", "Portfolio page", "Rehab & healthcare-AI expert page; LinkedIn; GitHub Pages; CV 2026 V3 refresh; later archived.", "portfolio/", "", "#40-#42"),
    ("2026-08-29", "Sheets", "4 Safety & Truth", "Google Sheets", "Fix", "Direct-first Sheets writes + import path", "Direct route, diagnostics, resilient writes.", "connectors/google_sheets_*", "", "#43-#45"),
    ("2026-08-29", "Direct Brief v2", "4 Safety & Truth", "Clinical / Rehab", "Enhancement", "Direct Brief v2 → v2.1", "Rebuilt /brief as direct Sheets runtime command; v2.1 adds Google Calendar.", "connectors/brief_runtime.py", "/brief", "#46-#49"),
    ("2026-08-29", "Gate 0.5", "4 Safety & Truth", "Security", "Security", "Fail closed on ambiguous Calendar dates/times", "Regression tests for ambiguity; invalid references (30 Aug).", "connectors/calendar_intent.py", "", "#63"),
    ("2026-08-30", "State Safety", "4 Safety & Truth", "Core State", "Enhancement", "StateStore concurrency + persistent markers", "Transactional Manager writes, loop markers, CI tests.", "engine/store.py, engine/manager.py", "", "#64"),
    ("2026-08-30", "WO-8", "4 Safety & Truth", "Inbox", "Feature", "Linked multi-intent capture", "Conservative multi-intent recorder; multi-intent-aware unified inbox.", "engine/multi_intent.py, connectors/multi_intent_runtime.py", "", "#65"),
    ("2026-08-30", "Super Manager v1.1", "4 Safety & Truth", "Chief of Staff", "Feature", "Super Manager", "Grounded in WO-8 record links; FAST compat with NEEDS_INPUT.", "connectors/super_manager.py", "/manager /manager_status", "#66"),
    ("2026-08-30", "Canary", "4 Safety & Truth", "Automation", "Feature", "FAST-only Manager canary (flag OFF)", "Wired into webhook runtime behind disabled flag.", "connectors/manager_fast_canary.py", "/manager_shadow", "#67"),
    ("2026-08-31", "Capability Truth", "4 Safety & Truth", "Governance", "Feature", "Capability Truth + action preflight", "Runtime truth about tools; privacy correction; blocks blanket text-only denial; shopping boundaries.", "connectors/capability_truth.py, capability_runtime.py", "/capabilities", "#68"),
    ("2026-08-31", "NL Actions", "4 Safety & Truth", "Actions", "Feature", "Natural-language Action Executor", "Approval receipts; deadline/reminder/project report; reminder negation; language safety.", "connectors/action_executor.py, action_runtime.py, action_deadline_report.py, action_language_safety.py", "/act /action_status /approve_action /reject_action", "#69"),
    ("2026-08-31", "Exec Brief v3", "4 Safety & Truth", "Clinical / Rehab", "Enhancement", "Executive signal discovery", "Constraints, logistics, commitments, status; StateStore route; 06:45 travel rule; departure calculation.", "connectors/executive_signals.py, brief_signal_runtime.py", "/brief", "#70"),
    # ---------------- 1-4 Sep ----------------
    ("2026-09-01", "Capability Truth v2", "5 Extensions", "Governance", "Enhancement", "Ground read-only Main Sheet requests", "Regressions for read-only sheet capability.", "connectors/capability_truth.py", "", "#71"),
    ("2026-09-01", "Commerce", "5 Extensions", "Commerce", "Feature", "Commerce Agent core", "Trusted checkout adapter, read-only deal scout, sandbox smoke test, Telegram runtime.", "connectors/commerce_agent.py, commerce_checkout.py, commerce_scout.py, commerce_sandbox.py, commerce_runtime.py", "/shop /prepare_order /approve_order /commerce_status /commerce_test", "#73"),
    ("2026-09-01", "Commerce", "5 Extensions", "Security", "Security", "$100 pilot order + daily caps", "Defense-in-depth cap.", "connectors/commerce_agent.py", "", "#73"),
    ("2026-09-02", "Bridge", "5 Extensions", "API", "Feature", "First-party /chat Bridge API", "BRIDGE_API_KEY-protected POST /chat; unified ask alias.", "connectors/bridge_api.py", "POST /chat", ""),
    ("2026-09-03", "CI", "5 Extensions", "Infra", "Infra", "Strategic shadow DEV runner", "Manual isolated runner; safe prerequisite reporting; bounded concurrency.", ".github/workflows/strategic-shadow-dev.yml", "", ""),
    ("2026-09-04", "Books v3.0", "5 Extensions", "Knowledge", "Feature", "Books context (learning shelf)", "Book retrieval, Super Manager integration, fast path in Telegram; v2.0 → v3.0.", "engine/books_context.py", "/books", ""),
    # ---------------- 8-9 Sep ----------------
    ("2026-09-08", "Connectors", "6 Master OS", "Google Docs", "Integration", "Google Docs connector + connection doctor", "Unified connection status checker; docs & tab creation from the bot.", "connectors/google_docs_service.py, connection_setup.py", "", "#78"),
    ("2026-09-08", "Connectors", "6 Master OS", "Docs", "Docs", "Interactive connection walkthrough page", "Checkboxes, print/PDF, offline; Buffer setup phase.", "docs/connection-walkthrough.html", "", "#78"),
    ("2026-09-09", "AI OS v0.9 Master OS", "6 Master OS", "Knowledge", "Feature", "Google Drive standard tree + drive_tree", "Abdulrahman_Master_OS roots 01–05; render/checklist.", "engine/master_os.py, drive_tree.py", "drive_tree.py render|checklist", "#80"),
    ("2026-09-09", "AI OS v0.9 Master OS", "6 Master OS", "Agents", "Feature", "Four sub-agents matrix", "Morning Briefing / Clinical & Ops / Knowledge & Audio / Finance & Life.", "prompts/master-os-agents.md", "", "#80"),
    ("2026-09-09", "AI OS v0.9 Master OS", "6 Master OS", "Automation", "Feature", "Automation scheduler (11 jobs, Riyadh)", "Daily 06:45/07:30/16:00/20:30; weekly Sun/Tue/Thu/Fri; monthly 28/1/last; drafts only.", "engine/scheduler.py", "/schedule, scheduler.py today-actions", "#80"),
    ("2026-09-09", "AI OS v0.9 Master OS", "6 Master OS", "Knowledge", "Feature", "Mind-map generator", "Markdown → Mermaid mindmap + text tree + library; Friday map.", "engine/mindmap.py", "/mind_maps", "#80"),
    ("2026-09-09", "AI OS v0.9 Master OS", "6 Master OS", "Voice", "Feature", "Audio digest pipeline", "QUEUED→DIGESTED→NARRATED; 5–7 min script; mp3 via ElevenLabs.", "engine/audio_digest.py", "/audio_digests", "#80"),
    ("2026-09-09", "AI OS v0.9 Master OS", "6 Master OS", "Telegram", "Command", "Master OS panel commands", "/masteros /schedule /mind_maps /audio_digests /today_actions /run", "engine/telegram_bot.py", "/masteros /run /today_actions", "#80"),
    ("2026-09-09", "v0.9 live wiring", "6 Master OS", "Voice", "Integration", "ElevenLabs voice channel", "Config + live probe in connection setup.", "connectors/connection_setup.py", "", "#81"),
    ("2026-09-09", "v0.9 live wiring", "6 Master OS", "Telegram", "Command", "/diag channel status", "Presence-only offline diagnostics; empty-state hardening.", "engine/telegram_bot.py", "/diag", "#81"),
    ("2026-09-09", "v0.9 live wiring", "6 Master OS", "Telegram", "Fix", "Wire Master OS commands to production runtime", "Commands & callbacks in webhook runtime.", "connectors/telegram_webhook_runtime.py", "", "#82"),
    ("2026-09-09", "v0.9 live wiring", "6 Master OS", "Telegram", "Fix", "Reply noise removed", "No routine footers; task-first brevity policy.", "connectors/telegram_webhook_runtime.py", "", "#83"),
    # ---------------- 11-12 Sep ----------------
    ("2026-09-11", "Proactive v1.0", "7 Proactive", "Proactive", "Feature", "Proactive Chief of Staff loop", "observe→remember→predict→score→decide→act/prepare/alert→learn in one Store transaction; from manager --loop.", "engine/proactive.py", "proactive.py sweep|brief|status|orders", "#85"),
    ("2026-09-11", "Proactive v1.0", "7 Proactive", "Proactive", "Feature", "Open loops register", "Commitments to self/from others, renewals, decisions, project risks.", "engine/proactive.py", "", "#85"),
    ("2026-09-11", "Proactive v1.0", "7 Proactive", "Governance", "Feature", "Autonomy ladder L0–L4 per category", "Comms/money/legal/health/reputation L0–L1; internal tasks L2–L3; no global free-act switch.", "engine/proactive.py", "", "#85"),
    ("2026-09-11", "Proactive v1.0", "7 Proactive", "Proactive", "Feature", "8 standing orders", "Meeting prep, appointment window, rotting waits, money due, renewals, conflicts, travel, missed-commitment recovery.", "engine/proactive.py", "order-disable SO-00x", "#85"),
    ("2026-09-11", "Proactive v1.0", "7 Proactive", "Governance", "Security", "Proactive guards", "Confidence<0.8 prepare only; 6 alerts/day; quiet hours 22:00–06:30; undo; pause; audit.", "engine/proactive.py", "undo PA-xxxx, pause --hours 4, feedback", "#85"),
    ("2026-09-11", "Proactive v1.0.1", "7 Proactive", "Telegram", "Integration", "Urgent-alert Telegram channel", "proactive_alert push + first live sweep.", "engine/proactive.py", "proactive.py push-test", "#85"),
    ("2026-09-11", "Proactive v1.0.2", "7 Proactive", "Telegram", "Command", "Proactive bot commands", "/proactive /sweep /proactive_test", "connectors/telegram_webhook_runtime.py", "/proactive /sweep /proactive_test", "#85"),
    ("2026-09-11", "Proactive v1.0.3", "7 Proactive", "Self-improvement", "Feature", "Weekly acceptance review + bounded self-tuning", "Adjusts confidence threshold / max alerts; guard ≥5 decisions.", "engine/proactive.py", "proactive.py review --apply", "#85"),
    ("2026-09-11", "Buffer", "7 Proactive", "Publishing", "Integration", "Buffer publisher", "GraphQL connector, setup guide, status check, --service resolution, GitHub Action, PR post queue.", "connectors/buffer_publisher.py, .github/workflows/buffer-publish.yml, posts/buffer/", "", "#90-#91"),
    ("2026-09-11", "Content", "7 Proactive", "Automation", "Feature", "Production proactive worker", "Runs in webhook every 15 min; last_proactive_worker marker.", "connectors/proactive_worker.py", "", "#95"),
    ("2026-09-11", "Content", "7 Proactive", "Publishing", "Feature", "Content Creator orchestra", "Researcher → critic → creator; preview; approve/reject; status.", "connectors/content_creator.py, content_runtime.py", "/content /approve_content /reject_content /content_status", "#95"),
    ("2026-09-11", "Content", "7 Proactive", "Google Sheets", "Integration", "Life Pulse publishing sheet", "Reads Ready/Idea rows; writes PUBLISH_QUEUE.", "connectors/content_sheet.py", "/content_sheet /content_sheet_status", "#96"),
    ("2026-09-11", "Content", "7 Proactive", "Publishing", "Fix", "Low OpenRouter credit resilience", "Workflow continues on low credit.", "connectors/content_creator.py", "", "#97"),
    ("2026-09-11", "Content", "7 Proactive", "Publishing", "Fix", "Truncated Content Creator JSON recovery", "Safe recovery.", "connectors/content_creator.py", "", "#98"),
    ("2026-09-11", "Content", "7 Proactive", "Governance", "Enhancement", "Runtime-truth grounding for publishing/proactive answers", "No invented status.", "connectors/capability_truth.py", "", "#99"),
    ("2026-09-12", "Content", "7 Proactive", "Publishing", "Feature", "Gemini image & video workflow", "Design/video content commands; Sheet queue IDs accepted.", "connectors/content_media.py", "/design_content /video_content /media_status", "#100-#101"),
    ("2026-09-12", "Books", "7 Proactive", "Knowledge", "Fix", "Learning-tab name variants", "Tolerate variants; stop silent fall-through.", "engine/books_context.py", "", "#102"),
    # ---------------- 14-15 Sep ----------------
    ("2026-09-14", "Prod fix", "8 Finance & Autopay", "Automation", "Fix", "Dispatch scheduler from production worker", "Due scheduler jobs run from webhook worker.", "connectors/proactive_worker.py", "", "#103"),
    ("2026-09-14", "Prod fix", "8 Finance & Autopay", "Memory", "Fix", "Scheduler persistence + Drive project memory", "Riyadh scheduler persisted; Google Drive project memory connector.", "connectors/project_memory.py", "", "#104"),
    ("2026-09-14", "Prod fix", "8 Finance & Autopay", "Runtime", "Fix", "Real-clock anchor + /time probe", "Stopped crash-loop; agent anchored to real clock.", "engine/runtime_clock.py", "/time", "#105"),
    ("2026-09-15", "AI OS v2.0", "8 Finance & Autopay", "Finance", "Feature", "Unified Finance Hub v2.0", "finance as single truth; webhook/CSV/manual pull; monthly snapshot with health_index; Thu 07:00 + 1st 07:30 + hourly.", "engine/finance_hub.py", "system_status.py --finance", "#106"),
    ("2026-09-15", "AI OS v2.0", "8 Finance & Autopay", "Automation", "Enhancement", "Dormant layers activated (14/14)", "asset_registry, backup_verify, change_intelligence, observability, trust_dashboard in manager loop.", "engine/manager.py", "system_status.py --layers --schedule", "#106"),
    ("2026-09-15", "AI OS v2.0", "8 Finance & Autopay", "Proactive", "Fix", "Proactive live context", "No more INFERENCE/MISSING for finance/proactive in bot.", "engine/system_status.py", "", "#106"),
    ("2026-09-15", "Web search", "8 Finance & Autopay", "Search", "Feature", "Verified YouTube search", "Only verified watch?v= links; auto-attach; DuckDuckGo default, optional YouTube API; PII stripped.", "connectors/web_search.py", "/youtube /search", "#107"),
    ("2026-09-15", "Money v1.1", "8 Finance & Autopay", "Finance", "Feature", "Money threshold autopay < 375 SAR", "money L1→L3 under threshold via payment gateway with receipt/undo; ≥375 or unknown stays red draft.", "connectors/payment_gateway.py, engine/proactive.py", "proactive.py threshold|matrix|status", "#108"),
    ("2026-09-15", "Money v1.1", "8 Finance & Autopay", "Security", "Security", "Autopay guards", "ENABLED flag + ACK phrase + gateway + known amount + daily cap 375 + idempotency; HARD_MAX_SAR in code; env can only lower.", "connectors/payment_gateway.py", "", "#108"),
    ("2026-09-15", "Money v1.1", "8 Finance & Autopay", "Core State", "Feature", "autopay_executions ledger + audit events", "autopay_authorized/executed/failed/blocked; sandbox connector; 18 tests; verify script.", "connectors/payment_sandbox.py, tests/test_money_threshold.py, scripts/verify_money_threshold.sh", "bash scripts/verify_money_threshold.sh", "#108"),
]

COMMANDS = [
    ("2026-08-22", "/start /help /profile /sources /selftest", "Bot basics, profile, source counts, self-test"),
    ("2026-08-22", "/okr /tasks /decisions /approve /door /mastery /reviews /answer", "OKR, tasks, decisions, approvals, life door, learning reviews"),
    ("2026-08-22", "/manager /manager_status /bedrock_test /context_test", "Manager loop status, model & context probes"),
    ("2026-08-24", "/brief", "Executive brief discovery (v1.3 → Direct Brief v2.1 → Exec Brief v3)"),
    ("2026-08-25", "/memory /memory_status /confirm_memory /update_memory /storage_status", "Provenance memory management"),
    ("2026-08-28", "/delegate /agents /council /mission /google_access", "Multi-agent delegation & missions"),
    ("2026-08-30", "/manager_shadow", "FAST manager canary (flag OFF)"),
    ("2026-08-31", "/capabilities /act /action_status /approve_action /reject_action", "Capability truth & natural-language actions"),
    ("2026-09-01", "/shop /prepare_order /approve_order /commerce_status /commerce_test", "Commerce agent ($100 pilot caps)"),
    ("2026-09-04", "/books", "Learning shelf book retrieval"),
    ("2026-09-09", "/masteros /schedule /mind_maps /mindmaps /digests /audio_digests /today_actions /run /diag", "Master OS panel & diagnostics"),
    ("2026-09-11", "/proactive /sweep /proactive_test", "Proactive Chief of Staff"),
    ("2026-09-11", "/content /approve_content /reject_content /content_status /content_sheet /content_sheet_status", "Content Creator & publishing sheet"),
    ("2026-09-12", "/design_content /video_content /media_status", "Gemini image/video content"),
    ("2026-09-14", "/time", "Real-clock liveness probe"),
    ("2026-09-15", "/youtube /search", "Verified YouTube search"),
]

INTEGRATIONS = [
    ("2026-08-21", "Google Drive (sheet import)", "engine/import_drive.py", "Active"),
    ("2026-08-21", "ChatGPT / VS Code bridge", "engine/export_for_chat.py", "Active"),
    ("2026-08-22", "Telegram (polling)", "connectors/telegram_bot.py", "Superseded by webhook 24 Aug"),
    ("2026-08-22", "AWS Bedrock (Claude)", "connectors/telegram_bot.py", "Active (fallback)"),
    ("2026-08-22", "AWS Transcribe", "connectors/aws_transcribe.py", "Active"),
    ("2026-08-22", "Google Sheets — Apps Script webhook", "connectors/google_sheets_webhook.gs", "Active"),
    ("2026-08-22", "Gmail / Calendar / Drive / GitHub live connectors", "connectors/google_workspace.py, github_live.py", "Active"),
    ("2026-08-22", "Railway (container + volume)", "Dockerfile", "Active"),
    ("2026-08-24", "Telegram (webhook)", "connectors/telegram_webhook.py", "Active"),
    ("2026-08-24", "Google Forms (supervisor, pre-visit)", "connectors/*.gs", "Active"),
    ("2026-08-25", "Google Calendar actions", "connectors/calendar_actions.py", "Active"),
    ("2026-08-27", "OpenRouter", "connectors/model_gateway.py", "Active (primary)"),
    ("2026-08-28", "OpenAI GPT & Google Gemini (direct)", "connectors/direct_specialists.py", "Active"),
    ("2026-08-28", "Bedrock Nova Micro (lean)", "connectors/bedrock_team.py", "Active"),
    ("2026-08-28", "Google Knowledge gateway", "—", "Rolled back (#21)"),
    ("2026-09-02", "Bridge API /chat", "connectors/bridge_api.py", "Active"),
    ("2026-09-08", "Google Docs", "connectors/google_docs_service.py", "Active"),
    ("2026-09-09", "ElevenLabs", "connectors/connection_setup.py, engine/audio_digest.py", "Active (key in env)"),
    ("2026-09-11", "Buffer (GraphQL)", "connectors/buffer_publisher.py", "Active"),
    ("2026-09-12", "Gemini image/video", "connectors/content_media.py", "Active"),
    ("2026-09-14", "Google Drive project memory", "connectors/project_memory.py", "Active"),
    ("2026-09-15", "YouTube Data API / DuckDuckGo", "connectors/web_search.py", "Active"),
    ("2026-09-15", "Payment gateway webhook (+sandbox)", "connectors/payment_gateway.py", "Active, opt-in guarded"),
]

DAILY = [
    ("2026-08-21", 13, 0, 0, 7, 3, 14), ("2026-08-22", 51, 8, 1, 13, 12, 40), ("2026-08-24", 52, 13, 3, 14, 13, 7),
    ("2026-08-25", 53, 14, 5, 14, 13, 14), ("2026-08-27", 53, 16, 6, 14, 13, 1), ("2026-08-28", 53, 26, 16, 14, 13, 68),
    ("2026-08-29", 53, 27, 16, 14, 13, 20), ("2026-08-30", 54, 30, 21, 14, 14, 30), ("2026-08-31", 54, 38, 24, 14, 14, 22),
    ("2026-09-01", 54, 43, 25, 14, 14, 13), ("2026-09-02", 54, 44, 25, 14, 14, 3), ("2026-09-03", 54, 44, 25, 14, 14, 6),
    ("2026-09-04", 55, 44, 25, 14, 14, 6), ("2026-09-08", 55, 46, 28, 14, 16, 4), ("2026-09-09", 60, 46, 33, 15, 18, 8),
    ("2026-09-11", 61, 51, 37, 16, 20, 13), ("2026-09-12", 61, 52, 39, 16, 20, 3), ("2026-09-14", 62, 54, 42, 16, 21, 3),
    ("2026-09-15", 64, 57, 44, 16, 29, 5),
]

HDR_FILL = PatternFill("solid", fgColor="1E293B")
HDR_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="CBD5E1")
TYPE_COLORS = {"Feature": "DCFCE7", "Enhancement": "DBEAFE", "Fix": "FEE2E2", "Integration": "E0F2FE", "Command": "F3E8FF",
               "Infra": "F1F5F9", "Security": "FEF3C7", "Docs": "F5F5F4", "Rollback": "FECACA", "Test": "ECFDF5"}


def sheet(wb, title, headers, rows, widths, table_name):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for r in rows:
        ws.append(list(r))
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(bottom=THIN)
    ws.freeze_panes = "A2"
    ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    t = Table(displayName=table_name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t)
    return ws


def main():
    wb = Workbook()
    wb.remove(wb.active)

    # --- Summary ---
    ws = wb.create_sheet("Summary")
    ws["A1"] = "Rehabilitation Chief-of-Staff Agent — Function Update Register"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = "Source: git history of main (304 commits / 108 PRs), README release notes, docs/v*.md · 21 Aug → 15 Sep 2026"
    ws["A2"].font = Font(italic=True, color="64748B")
    by_type = Counter(u[4] for u in U)
    by_phase = Counter(u[2] for u in U)
    ws["A4"], ws["B4"] = "Total updates", len(U)
    ws["A5"], ws["B5"] = "Telegram command groups", len(COMMANDS)
    ws["A6"], ws["B6"] = "Integrations", len(INTEGRATIONS)
    ws["A7"], ws["B7"] = "Distinct version lines", len({u[1] for u in U})
    ws["A9"], ws["B9"] = "By type", "Count"
    r = 10
    for k, v in by_type.most_common():
        ws.cell(r, 1, k); ws.cell(r, 2, v); r += 1
    type_end = r - 1
    r += 1
    ws.cell(r, 1, "By phase"); ws.cell(r, 2, "Count"); r += 1
    ph_start = r
    for k in sorted(by_phase):
        ws.cell(r, 1, k); ws.cell(r, 2, by_phase[k]); r += 1
    ph_end = r - 1
    for a in ("A4", "A5", "A6", "A7", "A9", "B9", f"A{ph_start-1}", f"B{ph_start-1}"):
        ws[a].font = Font(bold=True)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 12

    ch = BarChart(); ch.type = "bar"; ch.title = "Updates by type"; ch.height, ch.width = 8, 16
    ch.add_data(Reference(ws, min_col=2, min_row=9, max_row=type_end), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=10, max_row=type_end)); ch.legend = None
    ws.add_chart(ch, "D4")
    ch2 = BarChart(); ch2.title = "Updates by phase"; ch2.height, ch2.width = 8, 16
    ch2.add_data(Reference(ws, min_col=2, min_row=ph_start - 1, max_row=ph_end), titles_from_data=True)
    ch2.set_categories(Reference(ws, min_col=1, min_row=ph_start, max_row=ph_end)); ch2.legend = None
    ws.add_chart(ch2, "D22")

    # --- Function Updates ---
    headers = ["#", "Date", "Version line", "Phase", "Category", "Type", "Function / Feature", "Description", "Module / File", "Command / CLI", "PR"]
    rows = [(i + 1, *u) for i, u in enumerate(U)]
    wsu = sheet(wb, "Function Updates", headers, rows, [5, 12, 20, 18, 16, 13, 36, 70, 40, 34, 10], "Updates")
    for row in wsu.iter_rows(min_row=2):
        col = TYPE_COLORS.get(row[5].value)
        if col:
            row[5].fill = PatternFill("solid", fgColor=col)

    # --- Commands ---
    sheet(wb, "Telegram Commands", ["Introduced", "Commands", "Purpose"], COMMANDS, [12, 80, 50], "Commands")

    # --- Integrations ---
    sheet(wb, "Integrations", ["Introduced", "Integration", "Module", "Status"], INTEGRATIONS, [12, 46, 50, 28], "Integrations")

    # --- Growth ---
    g = sheet(wb, "Growth Metrics", ["Date", "Engine modules", "Connectors", "Test files", "Prompt packs", "Docs", "Commits/day"],
              DAILY, [12, 15, 12, 11, 13, 8, 13], "Growth")
    lc = LineChart(); lc.title = "Codebase growth"; lc.height, lc.width = 10, 22
    lc.add_data(Reference(g, min_col=2, max_col=6, min_row=1, max_row=len(DAILY) + 1), titles_from_data=True)
    lc.set_categories(Reference(g, min_col=1, min_row=2, max_row=len(DAILY) + 1))
    g.add_chart(lc, "I2")

    # --- Version lines ---
    lines = {}
    for u in U:
        lines.setdefault(u[1], [u[0], u[0], 0])
        lines[u[1]][1] = u[0]; lines[u[1]][2] += 1
    vrows = [(k, v[0], v[1], v[2]) for k, v in sorted(lines.items(), key=lambda kv: kv[1][0])]
    sheet(wb, "Version Lines", ["Version line", "First", "Last", "Updates"], vrows, [26, 12, 12, 10], "Versions")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    xlsx = OUT_DIR / "function-updates.xlsx"
    wb.save(xlsx)
    with (OUT_DIR / "function-updates.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(headers); w.writerows(rows)
    print(f"wrote {xlsx} ({len(U)} updates)")


if __name__ == "__main__":
    main()
